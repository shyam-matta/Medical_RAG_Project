"""
Generates the project report as a .docx file in the same style as the
sample RAG report (plain academic Word document, no heavy formatting).
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT = os.path.join(os.path.dirname(__file__), "Medical_RAG_Report.docx")

doc = Document()

# ── Page margins ────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1.2)
    section.right_margin  = Inches(1.2)

# ── Style helpers ───────────────────────────────────────────────────
def heading1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.size  = Pt(13)
    p.runs[0].font.bold  = True
    p.runs[0].font.color.rgb = RGBColor(0x1a, 0x56, 0xDB)
    return p

def heading2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.size  = Pt(11)
    p.runs[0].font.bold  = True
    p.runs[0].font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    return p

def para(text, bold_parts=None):
    """Add a normal paragraph. bold_parts = list of substrings to bold."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    if bold_parts:
        remaining = text
        for bp in bold_parts:
            idx = remaining.find(bp)
            if idx == -1:
                continue
            if idx > 0:
                run = p.add_run(remaining[:idx])
                run.font.size = Pt(11)
            bold_run = p.add_run(bp)
            bold_run.bold = True
            bold_run.font.size = Pt(11)
            remaining = remaining[idx + len(bp):]
        if remaining:
            run = p.add_run(remaining)
            run.font.size = Pt(11)
    else:
        run = p.add_run(text)
        run.font.size = Pt(11)
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.size = Pt(11)
        r2 = p.add_run(text[len(bold_prefix):])
        r2.font.size = Pt(11)
    else:
        r = p.add_run(text)
        r.font.size = Pt(11)
    return p

def numbered(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.size = Pt(11)
        r2 = p.add_run(text[len(bold_prefix):])
        r2.font.size = Pt(11)
    else:
        r = p.add_run(text)
        r.font.size = Pt(11)
    return p

def space():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)

# ── Title block ─────────────────────────────────────────────────────
title = doc.add_heading("Medical Information RAG Assistant", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.size  = Pt(16)
title.runs[0].font.bold  = True
title.runs[0].font.color.rgb = RGBColor(0x1a, 0x56, 0xDB)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Name: K. Kavur      Topic: Medical Information RAG System")
r.font.size = Pt(11)
r.font.bold = True

space()

# ════════════════════════════════════════════════════════
# PROBLEM STATEMENT
# ════════════════════════════════════════════════════════
heading1("Problem Statement")

para(
    "People frequently search the internet for health information, but general search results often "
    "mix reliable and unreliable sources. Asking a plain LLM health questions is also unreliable "
    "because the model may generate plausible-sounding but factually incorrect answers "
    "(hallucination), and it cannot cite the specific source it used. This project builds a "
    "Medical Information RAG Assistant where a user asks a general health question and the system "
    "answers only from a controlled set of trusted medical documents, along with the exact source "
    "document used for the answer. The system is strictly educational and does not diagnose "
    "patients or recommend personalised treatment."
)

space()

# ════════════════════════════════════════════════════════
# WHAT IS RAG AND WHY USE IT
# ════════════════════════════════════════════════════════
heading1("What is RAG and Why Use It")

para(
    "RAG stands for Retrieval Augmented Generation. Instead of relying on what the model already "
    "knows, the system first retrieves the most relevant parts of the knowledge base documents and "
    "then uses only that retrieved text to form the answer. This keeps the answer grounded in "
    "verified documents, eliminates hallucination since no generative model is used, avoids "
    "uncontrolled internet knowledge, and lets the system show the exact source document used for "
    "every answer, something a plain LLM call cannot do since it has no access to the local "
    "document collection."
)

space()

# ════════════════════════════════════════════════════════
# RAG PIPELINE
# ════════════════════════════════════════════════════════
heading1("RAG Pipeline")

para("The system follows two phases:")
space()

para("Indexing Phase (done once at startup):", bold_parts=["Indexing Phase (done once at startup):"])
para(
    "Document Loading  →  Text Chunking  →  TF-IDF Indexing  →  In-Memory Vector Store"
)
space()

para("Query Phase (done for every user question):", bold_parts=["Query Phase (done for every user question):"])
para(
    "User Question  →  TF-IDF Query Vector  →  Two-Stage Retrieval (top-k chunks)  →  "
    "Answer assembled from retrieved chunks  →  Answer with source citation and disclaimer"
)

space()

# ════════════════════════════════════════════════════════
# DETAILED WORKING
# ════════════════════════════════════════════════════════
heading1("Detailed Working of Each Stage")

heading2("Document Loading")
para(
    "Six plain-text medical documents are stored in a documents/ folder. On startup, each file is "
    "read and the source title is extracted from the first line of the document. The documents "
    "cover Diabetes, Hypertension, Common Cold, Nutrition Basics, Mental Health, and Exercise and "
    "Physical Health."
)

heading2("Chunking")
para(
    "Each document is split into overlapping word-level windows of 300 words with a 50-word "
    "overlap between consecutive chunks. The overlap ensures that ideas at chunk boundaries are "
    "not lost. Each chunk is tagged with its source document name and a unique chunk ID. "
    "Six documents produce 12 chunks in total."
)

heading2("TF-IDF Indexing")
para(
    "Each chunk is tokenized into lowercase words of length two or more. Term Frequency (TF) is "
    "computed per chunk as the normalised count of each term. Inverse Document Frequency (IDF) is "
    "computed across all chunks using the formula log((N+1)/(df+1)) + 1, where N is the total "
    "number of chunks and df is the number of chunks containing the term. This gives higher weight "
    "to rare, specific terms and lower weight to common words. The result is a sparse TF-IDF "
    "vector for every chunk stored in memory."
)

heading2("Two-Stage Retrieval")
para(
    "When the user asks a question, the same TF-IDF process converts the question into a query "
    "vector. Cosine similarity is then computed between the query vector and every chunk vector. "
    "To prevent false positives from coincidental word matches, a second keyword overlap filter "
    "is applied in parallel. A chunk must pass both conditions to be included in the results:"
)
bullet("Stage 1: TF-IDF cosine similarity score must be at least 0.15",
       bold_prefix="Stage 1:")
bullet("Stage 2: At least 25% of meaningful query words (length four or more) must appear in the chunk text",
       bold_prefix="Stage 2:")
para(
    "If no chunk passes both filters the system returns an out-of-scope response instead of "
    "guessing. The top-3 passing chunks are returned."
)

heading2("Answer Generation")
para(
    "The retrieved chunk texts are combined and returned directly as the answer body. No language "
    "model is used in this step, so hallucination is structurally impossible. The answer always "
    "includes the source document names, the retrieval scores for transparency, and a mandatory "
    "medical disclaimer stating that the information is educational only and not medical advice."
)

space()

# ════════════════════════════════════════════════════════
# TECHNOLOGIES USED
# ════════════════════════════════════════════════════════
heading1("Technologies Used")

numbered("Document Storage: Plain-text .txt files (no external database needed)",
         bold_prefix="Document Storage:")
numbered("Chunking: Custom word-window splitter with overlap (Python stdlib)",
         bold_prefix="Chunking:")
numbered("Retrieval: TF-IDF cosine similarity + keyword overlap ratio (Python stdlib: re, math)",
         bold_prefix="Retrieval:")
numbered("Answer Generation: Template-based, grounded assembly from retrieved chunks (no LLM)",
         bold_prefix="Answer Generation:")
numbered("Backend: Python 3.11 + Flask 3.0.3",
         bold_prefix="Backend:")
numbered("Frontend: Single-page HTML/CSS/JS chat interface served by Flask",
         bold_prefix="Frontend:")

space()

# ════════════════════════════════════════════════════════
# CHUNK SIZE, OVERLAP AND RETRIEVAL DEPTH
# ════════════════════════════════════════════════════════
heading1("Chunk Size, Overlap and Retrieval Depth")

para(
    "Chunk size used: 300 words with 50-word overlap (approximately 16% overlap). A smaller chunk "
    "loses surrounding context and a larger chunk pulls in irrelevant sentences from the same "
    "paragraph. The 50-word overlap ensures that a sentence spanning two chunks is fully captured "
    "by at least one of them."
)
para(
    "Retrieval depth (top-k) is set to 3 chunks per question. A lower k can miss information "
    "spread across multiple paragraphs, while a higher k introduces unrelated text that dilutes "
    "the answer. k=3 gave the best balance for the six-document knowledge base used here."
)
para(
    "The minimum cosine score threshold is 0.15 and the minimum keyword overlap is 0.25. These "
    "values were chosen by inspecting the actual scores of in-scope and out-of-scope test queries "
    "and picking thresholds that cleanly separated them."
)

space()

# ════════════════════════════════════════════════════════
# KNOWLEDGE BASE
# ════════════════════════════════════════════════════════
heading1("Knowledge Base")

para(
    "The knowledge base consists of six handcrafted educational medical documents. Each document "
    "is structured with clear section headings and factual content based on widely accepted health "
    "guidelines."
)

bullet("diabetes.txt: Type 1 and Type 2 diabetes, symptoms, blood sugar levels, complications, prevention",
       bold_prefix="diabetes.txt:")
bullet("hypertension.txt: Blood pressure categories, risk factors, lifestyle recommendations, monitoring",
       bold_prefix="hypertension.txt:")
bullet("common_cold.txt: Viral causes, symptoms, cold vs flu difference, home care, when to see a doctor",
       bold_prefix="common_cold.txt:")
bullet("nutrition_basics.txt: Macronutrients, vitamins and minerals, hydration, dietary patterns",
       bold_prefix="nutrition_basics.txt:")
bullet("mental_health.txt: Anxiety, depression, stress, wellness strategies, crisis resources",
       bold_prefix="mental_health.txt:")
bullet("exercise_health.txt: WHO exercise guidelines, types of exercise, benefits, sedentary behaviour",
       bold_prefix="exercise_health.txt:")

space()

# ════════════════════════════════════════════════════════
# VIVA POINTS
# ════════════════════════════════════════════════════════
heading1("Viva Points - Why RAG instead of directly asking an LLM")

numbered(
    "A plain LLM has no access to the local medical documents, so it cannot give source-cited "
    "answers grounded in those specific files.",
)
numbered(
    "LLMs can hallucinate, especially for specific medical facts. This system assembles answers "
    "only from retrieved document text, making hallucination structurally impossible."
)
numbered(
    "RAG lets the system show the exact source document used for each answer. A plain LLM "
    "response has no traceable source."
)
numbered(
    "The knowledge base is fully controlled. New documents can be added or removed without "
    "retraining or prompt engineering."
)
numbered(
    "The two-stage retrieval filter handles out-of-scope questions gracefully, returning a "
    "clear not-found message instead of a fabricated answer."
)

space()

# ════════════════════════════════════════════════════════
# WEB INTERFACE
# ════════════════════════════════════════════════════════
heading1("Web Interface")

para(
    "The system is accessible through a browser at http://127.0.0.1:5000. The interface is a "
    "single-page chat application served by Flask with the following features:"
)

bullet("Topic chips: Six clickable buttons for preset questions on each health topic",
       bold_prefix="Topic chips:")
bullet("Chat bubbles: User questions on the right in blue, assistant answers on the left in white",
       bold_prefix="Chat bubbles:")
bullet("Typing indicator: Animated dots shown while the server processes the query",
       bold_prefix="Typing indicator:")
bullet("Source panel: Document names and TF-IDF retrieval scores shown below every answer",
       bold_prefix="Source panel:")
bullet("Out-of-scope style: Yellow bubble with a clear message for unrecognised questions",
       bold_prefix="Out-of-scope style:")
bullet("Disclaimer box: Orange-bordered disclaimer on every assistant response",
       bold_prefix="Disclaimer box:")

space()
para("The backend exposes a single REST endpoint:", bold_parts=["The backend exposes a single REST endpoint:"])
para('POST /ask   Body: { "question": "..." }   Response: { "found", "answer", "sources", "scores" }')

space()

# ════════════════════════════════════════════════════════
# SAMPLE QUESTIONS
# ════════════════════════════════════════════════════════
heading1("Sample Questions the System Can Answer")

numbered(
    "What are the symptoms of diabetes? - answered from the diabetes.txt chunks covering symptoms section.",
)
numbered(
    "What is normal blood pressure? - answered from hypertension.txt blood pressure categories section."
)
numbered(
    "How do I treat a common cold at home? - answered from common_cold.txt general care section."
)
numbered(
    "What are macronutrients? - answered from nutrition_basics.txt macronutrients section."
)
numbered(
    "How does exercise affect mental health? - answered from mental_health.txt and exercise_health.txt chunks."
)
numbered(
    "What is the GDP of France? - system returns out-of-scope message since no matching chunk passes the filter."
)

space()

# ════════════════════════════════════════════════════════
# EVALUATION / TESTING
# ════════════════════════════════════════════════════════
heading1("Evaluation and Testing")

para(
    "The system was tested using an automated test script (test_rag.py) that runs ten test cases "
    "and checks whether the output matches the expected behaviour. All ten tests passed."
)

space()
para("In-scope queries tested (expected found = True):", bold_parts=["In-scope queries tested (expected found = True):"])
bullet("What are the symptoms of diabetes?  →  PASS")
bullet("What is normal blood pressure?  →  PASS")
bullet("How do I treat a common cold at home?  →  PASS")
bullet("What are macronutrients?  →  PASS")
bullet("How does exercise help mental health?  →  PASS")
bullet("What is the recommended weekly exercise amount?  →  PASS")

space()
para("Out-of-scope queries tested (expected found = False):", bold_parts=["Out-of-scope queries tested (expected found = False):"])
bullet("Tell me about rocket science  →  PASS")
bullet("Who won the football world cup?  →  PASS")
bullet("What is the GDP of France in 2024?  →  PASS")

space()
para("Additional checks:", bold_parts=["Additional checks:"])
bullet("Disclaimer present in every answer  →  PASS")

space()
para(
    "Answers were accurate when the relevant section was retrieved correctly. Wrong or incomplete "
    "answers mostly happened when the question wording did not closely match the document wording, "
    "which affected TF-IDF similarity scores. This confirms that retrieval quality is the main "
    "factor that determines final answer quality in a RAG system."
)

space()

# ════════════════════════════════════════════════════════
# LIMITATIONS
# ════════════════════════════════════════════════════════
heading1("Limitations")

numbered(
    "The retrieval method is TF-IDF which is keyword-based. It can miss semantically similar "
    "questions whose wording does not closely match the document text. Embedding-based semantic "
    "search would handle this better."
)
numbered(
    "The knowledge base is small (six documents, 12 chunks). Questions on health topics not "
    "covered by these documents will always return an out-of-scope response."
)
numbered(
    "Answer generation is template-based. The system returns the raw chunk text as the answer, "
    "which is sometimes verbose or starts mid-sentence at a chunk boundary."
)
numbered(
    "There is no conversation memory. Each question is answered independently with no awareness "
    "of previous questions in the same session."
)
numbered(
    "The system handles text only. Images, tables and diagrams in medical documents cannot be "
    "processed."
)

space()

# ════════════════════════════════════════════════════════
# FUTURE SCOPE
# ════════════════════════════════════════════════════════
heading1("Future Scope")

numbered(
    "Replace TF-IDF with sentence-transformer embeddings for semantic retrieval that handles "
    "paraphrased questions correctly."
)
numbered(
    "Add a language model generation step so answers are summarised and readable rather than "
    "raw chunk text, while still being grounded in retrieved content."
)
numbered(
    "Expand the knowledge base with more medical topics and allow users to upload their own "
    "trusted documents."
)
numbered(
    "Add hybrid search combining keyword (TF-IDF) and semantic (embedding) similarity for "
    "better retrieval accuracy."
)
numbered(
    "Add conversation memory so follow-up questions can reference the previous answer context."
)
numbered(
    "Add a re-ranking step after initial retrieval to select the best chunks before assembling "
    "the answer."
)

space()

# ════════════════════════════════════════════════════════
# CONCLUSION
# ════════════════════════════════════════════════════════
heading1("Conclusion")

para(
    "The Medical Information RAG Assistant successfully answers general health questions on "
    "Diabetes, Hypertension, Common Cold, Nutrition, Mental Health, and Exercise by combining "
    "document chunking, TF-IDF indexing, two-stage retrieval and grounded answer assembly. "
    "Compared to directly asking an LLM, this RAG-based approach gives more accurate, verifiable "
    "and source-cited answers. The two-stage filter prevents the system from answering questions "
    "outside the knowledge base, and the mandatory disclaimer ensures the educational-only purpose "
    "is always communicated to the user. The tuning of chunk size, overlap and retrieval depth "
    "directly impacts answer quality, and retrieval precision is the single most important factor "
    "in the overall system performance."
)

# Save
doc.save(OUTPUT)
print(f"Report saved: {OUTPUT}")
